"""
Generate a test Excel file for parsing validation.
Run this on the bastion after installing openpyxl:
    pip3 install openpyxl
    python3 scripts/create_test_excel.py
"""
from openpyxl import Workbook

wb = Workbook()

# Sheet 1: Team Members
ws1 = wb.active
ws1.title = "Team Members"
ws1.append(["Name", "Role", "Team", "Focus Area", "Start Date"])
ws1.append(["Sony Genu", "Sr. Manager", "RAG Ingestion", "AI Strategy", "2020-03-15"])
ws1.append(["Alice Chen", "Principal Engineer", "RAG Ingestion", "Architecture", "2019-07-01"])
ws1.append(["Bob Martinez", "Sr. SDE", "RAG Ingestion", "NLP/Extraction", "2021-01-10"])
ws1.append(["Carol Zhang", "SDE II", "RAG Ingestion", "Graph Database", "2022-06-20"])
ws1.append(["Dave Wilson", "SDE I", "RAG Ingestion", "Document Parsing", "2023-09-05"])

# Sheet 2: Services
ws2 = wb.create_sheet("Services")
ws2.append(["Service Name", "Owner", "Language", "Dependencies", "SLA (p99)"])
ws2.append(["Plato Ingestion", "Alice Chen", "Java/Python", "S3, SQS, Bedrock", "500ms"])
ws2.append(["Document Parser", "Dave Wilson", "Python", "Docling, S3", "2s"])
ws2.append(["Entity Extractor", "Bob Martinez", "Python", "Bedrock, Neptune", "5s"])
ws2.append(["Graph Loader", "Carol Zhang", "Python", "Neptune, OpenSearch", "3s"])
ws2.append(["Query Engine", "Alice Chen", "Python", "Neptune, Bedrock", "1s"])

# Sheet 3: On-Call Schedule
ws3 = wb.create_sheet("On-Call Schedule")
ws3.append(["Week Of", "Primary", "Secondary", "Escalation"])
ws3.append(["Jun 9, 2026", "Carol Zhang", "Dave Wilson", "Sony Genu"])
ws3.append(["Jun 16, 2026", "Bob Martinez", "Alice Chen", "Sony Genu"])
ws3.append(["Jun 23, 2026", "Dave Wilson", "Carol Zhang", "Sony Genu"])
ws3.append(["Jun 30, 2026", "Alice Chen", "Bob Martinez", "Sony Genu"])

output_path = "data/sample-team-data.xlsx"
wb.save(output_path)
print(f"✓ Created test Excel file: {output_path}")
print(f"  Sheets: {wb.sheetnames}")
print(f"  Sheet 1 (Team Members): 5 rows")
print(f"  Sheet 2 (Services): 5 rows")
print(f"  Sheet 3 (On-Call Schedule): 4 rows")
